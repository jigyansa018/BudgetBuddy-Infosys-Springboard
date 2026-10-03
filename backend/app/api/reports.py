from io import BytesIO
from datetime import date
import calendar

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from ..database import get_db
from ..auth import get_current_user
from ..models.user import User
from ..models.income import Income
from ..models.expense import Expense
from ..models.savings import Savings
from ..models.budget import Budget


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)


def get_month_dates(month: int, year: int):
    last_day = calendar.monthrange(year, month)[1]

    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)

    return start_date, end_date


def get_report_data(
    db: Session,
    current_user: User,
    month: int,
    year: int
):
    start_date, end_date = get_month_dates(month, year)

    # -------------------------
    # Income
    # -------------------------
    total_income = (
        db.query(func.coalesce(func.sum(Income.amount), 0))
        .filter(
            Income.user_id == current_user.id,
            Income.income_date >= start_date,
            Income.income_date <= end_date
        )
        .scalar()
    )

    # -------------------------
    # Expenses
    # -------------------------
    total_expenses = (
        db.query(func.coalesce(func.sum(Expense.amount), 0))
        .filter(
            Expense.user_id == current_user.id,
            Expense.expense_date >= start_date,
            Expense.expense_date <= end_date
        )
        .scalar()
    )

    total_income = float(total_income or 0)
    total_expenses = float(total_expenses or 0)

    remaining_balance = total_income - total_expenses

    # -------------------------
    # Expenses by category
    # -------------------------
    category_results = (
        db.query(
            Expense.category,
            func.coalesce(func.sum(Expense.amount), 0).label("total")
        )
        .filter(
            Expense.user_id == current_user.id,
            Expense.expense_date >= start_date,
            Expense.expense_date <= end_date
        )
        .group_by(Expense.category)
        .order_by(func.sum(Expense.amount).desc())
        .all()
    )

    categories = [
        {
            "category": category or "Uncategorized",
            "total": float(total or 0)
        }
        for category, total in category_results
    ]

    # -------------------------
    # Savings goals
    # -------------------------
    savings_goals = (
        db.query(Savings)
        .filter(Savings.user_id == current_user.id)
        .order_by(Savings.id.desc())
        .all()
    )

    savings = []

    for goal in savings_goals:
        target = float(goal.target_amount or 0)
        current = float(goal.current_amount or 0)

        remaining = max(target - current, 0)

        progress = (
            (current / target) * 100
            if target > 0
            else 0
        )

        savings.append({
            "name": goal.name,
            "target_amount": target,
            "current_amount": current,
            "remaining_amount": remaining,
            "progress_percentage": round(min(progress, 100), 2),
            "status": goal.status
        })

    # -------------------------
    # Budgets
    # -------------------------
    budgets = (
        db.query(Budget)
        .filter(
            Budget.user_id == current_user.id,
            Budget.month == month,
            Budget.year == year
        )
        .order_by(Budget.id)
        .all()
    )

    budget_data = []

    for budget in budgets:

        expense_query = db.query(
            func.coalesce(func.sum(Expense.amount), 0)
        ).filter(
            Expense.user_id == current_user.id,
            Expense.expense_date >= start_date,
            Expense.expense_date <= end_date
        )

        if budget.category:
            expense_query = expense_query.filter(
                Expense.category == budget.category
            )

        spent = float(expense_query.scalar() or 0)
        planned = float(budget.planned_amount or 0)

        remaining = planned - spent

        utilization = (
            (spent / planned) * 100
            if planned > 0
            else 0
        )

        budget_data.append({
            "category": budget.category or "Overall",
            "planned_amount": planned,
            "spent_amount": spent,
            "remaining_amount": remaining,
            "utilization_percentage": round(
                utilization,
                2
            )
        })

    return {
        "month": month,
        "year": year,
        "start_date": start_date,
        "end_date": end_date,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "remaining_balance": remaining_balance,
        "categories": categories,
        "savings": savings,
        "budgets": budget_data
    }


# =========================================================
# PDF REPORT
# =========================================================

@router.get("/monthly/pdf")
def monthly_pdf_report(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2000, le=2100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    data = get_report_data(
        db,
        current_user,
        month,
        year
    )

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    story = []

    # Title
    story.append(
        Paragraph(
            "BudgetBuddy Monthly Financial Report",
            styles["Title"]
        )
    )

    story.append(Spacer(1, 10))

    story.append(
        Paragraph(
            f"Report Period: "
            f"{calendar.month_name[month]} {year}",
            styles["Heading2"]
        )
    )

    story.append(
        Paragraph(
            f"User: {current_user.full_name}",
            styles["Normal"]
        )
    )

    story.append(Spacer(1, 20))

    # -------------------------
    # Financial Summary
    # -------------------------

    story.append(
        Paragraph(
            "Financial Summary",
            styles["Heading2"]
        )
    )

    summary_data = [
        ["Metric", "Amount"],
        [
            "Total Income",
            f"₹{data['total_income']:.2f}"
        ],
        [
            "Total Expenses",
            f"₹{data['total_expenses']:.2f}"
        ],
        [
            "Remaining Balance",
            f"₹{data['remaining_balance']:.2f}"
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[250, 150]
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "ALIGN",
                (1, 1),
                (-1, -1),
                "RIGHT"
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    story.append(summary_table)

    story.append(Spacer(1, 20))

    # -------------------------
    # Expense Categories
    # -------------------------

    story.append(
        Paragraph(
            "Expenses by Category",
            styles["Heading2"]
        )
    )

    category_table_data = [
        ["Category", "Amount"]
    ]

    for item in data["categories"]:
        category_table_data.append([
            item["category"],
            f"₹{item['total']:.2f}"
        ])

    if len(category_table_data) == 1:
        category_table_data.append([
            "No expenses",
            "₹0.00"
        ])

    category_table = Table(
        category_table_data,
        colWidths=[250, 150]
    )

    category_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "ALIGN",
                (1, 1),
                (-1, -1),
                "RIGHT"
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    story.append(category_table)

    story.append(Spacer(1, 20))

    # -------------------------
    # Savings Goals
    # -------------------------

    story.append(
        Paragraph(
            "Savings Goals",
            styles["Heading2"]
        )
    )

    savings_table_data = [
        [
            "Goal",
            "Target",
            "Saved",
            "Remaining",
            "Progress"
        ]
    ]

    for item in data["savings"]:
        savings_table_data.append([
            item["name"],
            f"₹{item['target_amount']:.2f}",
            f"₹{item['current_amount']:.2f}",
            f"₹{item['remaining_amount']:.2f}",
            f"{item['progress_percentage']:.1f}%"
        ])

    if len(savings_table_data) == 1:
        savings_table_data.append([
            "No savings goals",
            "-",
            "-",
            "-",
            "-"
        ])

    savings_table = Table(
        savings_table_data,
        colWidths=[130, 90, 90, 90, 70]
    )

    savings_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "ALIGN",
                (1, 1),
                (-1, -1),
                "RIGHT"
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    story.append(savings_table)

    story.append(Spacer(1, 20))

    # -------------------------
    # Budgets
    # -------------------------

    story.append(
        Paragraph(
            "Budget Utilization",
            styles["Heading2"]
        )
    )

    budget_table_data = [
        [
            "Category",
            "Planned",
            "Spent",
            "Remaining",
            "Used"
        ]
    ]

    for item in data["budgets"]:
        budget_table_data.append([
            item["category"],
            f"₹{item['planned_amount']:.2f}",
            f"₹{item['spent_amount']:.2f}",
            f"₹{item['remaining_amount']:.2f}",
            f"{item['utilization_percentage']:.1f}%"
        ])

    if len(budget_table_data) == 1:
        budget_table_data.append([
            "No budgets",
            "-",
            "-",
            "-",
            "-"
        ])

    budget_table = Table(
        budget_table_data,
        colWidths=[120, 90, 90, 90, 70]
    )

    budget_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "ALIGN",
                (1, 1),
                (-1, -1),
                "RIGHT"
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    story.append(budget_table)

    story.append(Spacer(1, 20))

    story.append(
        Paragraph(
            "Generated by BudgetBuddy",
            styles["Normal"]
        )
    )

    document.build(story)

    buffer.seek(0)

    filename = (
        f"BudgetBuddy_Report_"
        f"{year}_{month:02d}.pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )


# =========================================================
# EXCEL REPORT
# =========================================================

@router.get("/monthly/excel")
def monthly_excel_report(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2000, le=2100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):

    data = get_report_data(
        db,
        current_user,
        month,
        year
    )

    workbook = Workbook()

    # =====================================================
    # Summary Sheet
    # =====================================================

    summary = workbook.active
    summary.title = "Financial Summary"

    summary.append([
        "BudgetBuddy Monthly Financial Report"
    ])

    summary.append([
        "User",
        current_user.full_name
    ])

    summary.append([
        "Period",
        f"{calendar.month_name[month]} {year}"
    ])

    summary.append([])

    summary.append([
        "Metric",
        "Amount"
    ])

    summary.append([
        "Total Income",
        data["total_income"]
    ])

    summary.append([
        "Total Expenses",
        data["total_expenses"]
    ])

    summary.append([
        "Remaining Balance",
        data["remaining_balance"]
    ])

    # Formatting
    summary["A1"].font = Font(
        bold=True,
        size=16
    )

    for cell in summary[5]:
        cell.font = Font(bold=True)

    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 25

    # =====================================================
    # Expense Categories Sheet
    # =====================================================

    categories_sheet = workbook.create_sheet(
        "Expenses by Category"
    )

    categories_sheet.append([
        "Category",
        "Amount"
    ])

    for item in data["categories"]:
        categories_sheet.append([
            item["category"],
            item["total"]
        ])

    for cell in categories_sheet[1]:
        cell.font = Font(bold=True)

    categories_sheet.column_dimensions["A"].width = 25
    categories_sheet.column_dimensions["B"].width = 20

    # =====================================================
    # Savings Sheet
    # =====================================================

    savings_sheet = workbook.create_sheet(
        "Savings Goals"
    )

    savings_sheet.append([
        "Goal",
        "Target Amount",
        "Current Amount",
        "Remaining",
        "Progress %",
        "Status"
    ])

    for item in data["savings"]:
        savings_sheet.append([
            item["name"],
            item["target_amount"],
            item["current_amount"],
            item["remaining_amount"],
            item["progress_percentage"],
            item["status"]
        ])

    for cell in savings_sheet[1]:
        cell.font = Font(bold=True)

    savings_sheet.column_dimensions["A"].width = 25
    savings_sheet.column_dimensions["B"].width = 18
    savings_sheet.column_dimensions["C"].width = 18
    savings_sheet.column_dimensions["D"].width = 18
    savings_sheet.column_dimensions["E"].width = 15
    savings_sheet.column_dimensions["F"].width = 15

    # =====================================================
    # Budget Sheet
    # =====================================================

    budget_sheet = workbook.create_sheet(
        "Budget Utilization"
    )

    budget_sheet.append([
        "Category",
        "Planned Amount",
        "Spent Amount",
        "Remaining Amount",
        "Utilization %"
    ])

    for item in data["budgets"]:
        budget_sheet.append([
            item["category"],
            item["planned_amount"],
            item["spent_amount"],
            item["remaining_amount"],
            item["utilization_percentage"]
        ])

    for cell in budget_sheet[1]:
        cell.font = Font(bold=True)

    budget_sheet.column_dimensions["A"].width = 25
    budget_sheet.column_dimensions["B"].width = 20
    budget_sheet.column_dimensions["C"].width = 20
    budget_sheet.column_dimensions["D"].width = 20
    budget_sheet.column_dimensions["E"].width = 20

    # Alignment
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

    # Save workbook to memory
    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    filename = (
        f"BudgetBuddy_Report_"
        f"{year}_{month:02d}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        }
    )